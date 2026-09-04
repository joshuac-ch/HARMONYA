
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import random
from datetime import date, datetime, timedelta

# ============================================================
# CONFIGURACIÓN
# ============================================================
CANTIDAD = 50
ARCHIVO_SALIDA = "leads_hackathon_lead_fifty.xlsx"

# Fecha de referencia para generar datos ficticios
FECHA_REFERENCIA = datetime(2026, 8, 28, 23, 59, 59)

# Generar datos reproducibles
random.seed(20260828)

# ============================================================
# CATÁLOGOS DE DATOS FICTICIOS
# ============================================================
nombres_m = [
    "alejandro", "andres", "bruno", "carlos", "daniel", "diego",
    "eduardo", "fernando", "gabriel", "gustavo", "javier", "jorge",
    "juan", "lucas", "marco", "mateo", "miguel", "nicolas",
    "oscar", "rafael", "roberto", "sebastian", "sergio", "tomas"
]

nombres_f = [
    "andrea", "camila", "carla", "diana", "elena", "gabriela",
    "isabela", "karla", "laura", "lucia", "maria", "natalia",
    "paola", "patricia", "renata", "sofia", "valentina", "vanessa"
]

apellidos = [
    "garcia", "flores", "quispe", "condori", "mendoza", "torres",
    "ramirez", "castillo", "huaman", "rojas", "fernandez", "perez",
    "sanchez", "diaz", "rivera", "vargas", "mamani", "salazar",
    "chavez", "espinoza", "medina", "ortiz", "herrera", "ramos"
]

departamentos = {
    "arequipa": {
        "provincias": ["arequipa", "caylloma", "camana"],
        "distritos": ["paucarpata", "yanahuara", "cerro colorado", "miraflores"]
    },
    "cusco": {
        "provincias": ["cusco", "urubamba", "calca"],
        "distritos": ["cusco", "wanchaq", "san jeronimo", "san sebastian"]
    },
    "lima": {
        "provincias": ["lima", "callao", "huaral"],
        "distritos": ["miraflores", "surco", "san miguel", "los olivos"]
    },
    "la libertad": {
        "provincias": ["trujillo", "ascope", "pacasmayo"],
        "distritos": ["trujillo", "victor larco", "la esperanza", "huanchaco"]
    },
    "piura": {
        "provincias": ["piura", "sullana", "talara"],
        "distritos": ["piura", "castilla", "veintiseis de octubre", "sullana"]
    }
}

proyectos = [
    "aurora", "santorini", "alameda", "los olivos",
    "valle verde", "mirador", "las palmeras"
]

agentes = [
    "anderson", "mijael", "gabriel", "valentina", "lucas"
]

canales = [
    "digital", "referido", "presencial", "organico"
]

campanas = [
    "ninguna", "preventa", "verano", "vivienda2026", "blackweek"
]

medios = [
    "chatbot", "facebook", "instagram", "google", "whatsapp"
]

calles = [
    "jose mareategui", "victor lira", "los incas",
    "avenida ejercito", "avenida arequipa", "los sauces",
    "las flores", "miguel grau", "bolognesi"
]

etapas= [
  "Bandeja",
  "Contacto sin respuesta",
  "Seguimiento",
  "Lead tracking",
  "Desestimado",
] 

# ============================================================
# GENERADORES ÚNICOS
# ============================================================
dnis_usados = set()
telefonos_usados = set()
usuarios_usados = set()
correos_usados = set()


def generar_dni():
    """Genera un DNI ficticio de 8 dígitos sin repetir."""
    while True:
        dni = str(random.randint(10_000_000, 99_999_999))

        if dni not in dnis_usados:
            dnis_usados.add(dni)
            return dni


def generar_celular():
    """Genera un celular peruano ficticio de 9 dígitos sin repetir."""
    while True:
        celular = "9" + str(random.randint(10_000_000, 99_999_999))

        if celular not in telefonos_usados:
            telefonos_usados.add(celular)
            return celular


def generar_usuario_whatsapp(nombre, apellido):
    """Genera un usuario de WhatsApp ficticio y único."""
    base = f"{nombre}{apellido}".replace(" ", "")
    contador = 1

    while True:
        usuario = (
            base
            if contador == 1
            else f"{base}{contador}"
        )

        if usuario not in usuarios_usados:
            usuarios_usados.add(usuario)
            return usuario

        contador += 1


def generar_correo(nombre, apellido):
    """
    Genera un correo ficticio y único.
    Se usa example.com para evitar utilizar correos reales.
    """
    base = f"{nombre}.{apellido}".replace(" ", "")
    contador = 1

    while True:
        correo = (
            f"{base}@example.com"
            if contador == 1
            else f"{base}{contador}@example.com"
        )

        if correo not in correos_usados:
            correos_usados.add(correo)
            return correo

        contador += 1


def generar_fecha_nacimiento():
    """Genera una edad ficticia entre 21 y 65 años."""
    dias_min = 21 * 365
    dias_max = 65 * 365

    return (
        FECHA_REFERENCIA.date()
        - timedelta(days=random.randint(dias_min, dias_max))
    )


def generar_fecha_ingreso():
    """
    Genera una fecha y hora ficticia de llegada del lead.

    Los leads pueden haber llegado durante los últimos 30 días.
    Horario comercial aproximado: 08:00 - 20:00.
    """

    dias_atras = random.randint(0, 30)

    fecha = (
        FECHA_REFERENCIA
        - timedelta(days=dias_atras)
    )

    hora = random.randint(8, 20)
    minuto = random.randint(0, 59)
    segundo = random.randint(0, 59)

    return fecha.replace(
        hour=hora,
        minute=minuto,
        second=segundo
    )


# ============================================================
# CREAR EXCEL
# ============================================================
columnas = [
    "nombre",
    "apellido",
    "tipo_documento",
    "numero",
    "genero",
    "fecha_nacimiento",
    "pais_origen",
    "canal_origen",
    "agente_asignado",
    "nombre_campaña",
    "medio_captacion",
    "celular",
    "usuario_wasap",
    "correo_electronico",
    "departamento",
    "provincia",
    "distrito",
    "direccion",
    "proyecto",
    "fecha_ingreso",
    "hora_ingreso",
    "etapa"
]

wb = Workbook()
ws = wb.active
ws.title = "leads"

# ============================================================
# ENCABEZADOS
# ============================================================
for col, nombre_columna in enumerate(columnas, start=1):

    celda = ws.cell(
        row=1,
        column=col,
        value=nombre_columna
    )

    celda.font = Font(bold=True)

    celda.fill = PatternFill(
        "solid",
        fgColor="D9EAF7"
    )

    celda.alignment = Alignment(
        horizontal="center"
    )


# ============================================================
# GENERAR REGISTROS
# ============================================================
for fila in range(2, CANTIDAD + 2):

    genero = random.choice([
        "masculino",
        "femenino"
    ])

    nombre = (
        random.choice(nombres_m)
        if genero == "masculino"
        else random.choice(nombres_f)
    )

    apellido = random.choice(apellidos)

    departamento = random.choice(
        list(departamentos.keys())
    )

    ubicacion = departamentos[departamento]

    provincia = random.choice(
        ubicacion["provincias"]
    )

    distrito = random.choice(
        ubicacion["distritos"]
    )

    dni = generar_dni()
    celular = generar_celular()

    usuario_wasap = generar_usuario_whatsapp(
        nombre,
        apellido
    )

    correo = generar_correo(
        nombre,
        apellido
    )

    # Fecha y hora de llegada del lead
    fecha_ingreso = generar_fecha_ingreso()

    datos = [
        nombre,
        apellido,
        "DNI",
        dni,
        genero,
        generar_fecha_nacimiento(),
        "peru",
        random.choice(canales),
        random.choice(agentes),
        random.choice(campanas),
        random.choice(medios),
        celular,
        usuario_wasap,
        correo,
        departamento,
        provincia,
        distrito,
        random.choice(calles),
        random.choice(proyectos),

        # Fecha de llegada
        fecha_ingreso.date(),

        # Hora de llegada
        fecha_ingreso.time(),
        
        random.choice(etapas),
    ]

    for col, valor in enumerate(datos, start=1):

        celda = ws.cell(
            row=fila,
            column=col,
            value=valor
        )

        # Fecha de nacimiento
        if col == 6:
            celda.number_format = "DD/MM/YYYY"

        # DNI y celular como texto
        if col in (4, 12):
            celda.number_format = "@"

        # Fecha de ingreso
        if col == 20:
            celda.number_format = "DD/MM/YYYY"

        # Hora de ingreso
        if col == 21:
            celda.number_format = "HH:MM:SS"


# ============================================================
# FORMATO DEL EXCEL
# ============================================================
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

anchos = {
    "A": 14,
    "B": 16,
    "C": 16,
    "D": 14,
    "E": 14,
    "F": 20,
    "G": 14,
    "H": 16,
    "I": 18,
    "J": 20,
    "K": 18,
    "L": 15,
    "M": 22,
    "N": 32,
    "O": 18,
    "P": 18,
    "Q": 20,
    "R": 24,
    "S": 20,
    "T": 18,
    "U": 16,
    "V": 14
}

for columna, ancho in anchos.items():
    ws.column_dimensions[columna].width = ancho


# ============================================================
# GUARDAR
# ============================================================
wb.save(ARCHIVO_SALIDA)

print(f"Excel generado: {ARCHIVO_SALIDA}")
print(f"Registros generados: {CANTIDAD}")
print(f"DNIs únicos: {len(dnis_usados)}")
print(f"Celulares únicos: {len(telefonos_usados)}")
print(f"Usuarios WhatsApp únicos: {len(usuarios_usados)}")
print(f"Correos únicos: {len(correos_usados)}")
